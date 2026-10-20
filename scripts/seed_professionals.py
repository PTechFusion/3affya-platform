"""Imports the doctors of the NPPES registry shortlist into public.professionals as unclaimed listings (task 1.5, L1).

The shortlist is the October 2026 NPPES release scored for likely Arabic-speaking professionals
(aafia_registry_shortlist_2026-10.xlsx, built outside this repo). Only doctors get a card and a profile (decided 20 Oct
2026): the three doctor sheets "Physicians High", "Physicians Medium" and "Dentists" (28,587 rows) are the import.
"Other Professions" (students and every non-doctor profession) and "Low (review)" (rows that need a manual check) are
refused by name, and so is any other sheet. A second guard works row by row: a row is imported only when its Profession
is Physician with a physician's taxonomy code (NUCC grouping Allopathic & Osteopathic Physicians, codes starting 20) or
Dentist with a dentist's code (NUCC classification Dentist: codes starting 1223, and 125Q00000X Oral Medicine). Any other
row of the doctor sheets is counted and reported as not imported, with the reason, and never written (the October file
has three: Denturist codes under the Dentist profession).

What a row carries: the registry names, credential, profession, taxonomy code, practice city and US state, the practice
country (ISO 3166-1 alpha-2: 'US' when the State is a US state, DC or territory code; for a doctor practising abroad, the
country the registry's State, or State and City, names unambiguously, from the short explicit table below; null where
it names none, as for the US military mail codes AA, AE and AP), a display name, and the specialty group its taxonomy
code maps to (public.specialties.taxonomy_codes, written by map_taxonomy.py; an unmapped code leaves specialty_id null,
never a default). What it does not carry: address, phone, fax, emails, old address, the moved flag, the registry's
last-updated date or the outreach scores. An unclaimed page shows no contact details (ADR-5) and outreach emails are a
separate decision. registry_source and registry_as_of describe this NPPES row; the CMS enrichment (enrich_cms.py) keeps
its own date in cms_as_of, so neither script overwrites the other's provenance.

Idempotent on NPI. New NPIs are inserted; an existing row changes only while it is still unclaimed and from nppes AND an
imported value or the registry date differs (a newer --registry-as-of refreshes registry_as_of on a row whose facts are
unchanged, so the date always names the release the row was last read from), so a second run with the same file changes
zero rows. A claimed (or removed, or self-listed) row is never touched, nor are its slug and specialty links. The slug is
set once by the database trigger on insert (firstname-lastname-city-npiSuffix) and never rewritten here. The script never
deletes: a row an earlier import wrote outside the doctor scope stays until someone decides to remove it. Everything runs
in one transaction; --dry-run rolls it back. Done when the rows imported equal the doctor rows of the shortlist that pass
the guard (exit code 1 otherwise).

Usage:
  python scripts/seed_professionals.py --shortlist PATH --registry-as-of 2026-10-01 [--sheets ...] [--dry-run]

Report: stdout and output/seed_professionals_report.md (real runs only). The report file holds aggregate counts only
(rows per sheet, reason, profession, taxonomy code and country); the NPIs of rejected and not imported rows and the raw
State values abroad are printed to stdout, never written to it.
"""

from __future__ import annotations

import argparse
import collections
import csv
import dataclasses
import datetime as dt
import hashlib
import io
import pathlib
import re
import sys
import time
from typing import Sequence

# Importable under `python -I` too (isolated mode leaves the script's folder off sys.path).
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from seed_common import OUTPUT_DIR, add_database_argument, clean, connect, markdown_table, utc_now, write_report  # noqa: E402

# The doctor sheets: the only sheets imported (decided 20 Oct 2026). map_taxonomy.py reads the same three.
DOCTOR_SHEETS = ("Physicians High", "Physicians Medium", "Dentists")
DEFAULT_SHEETS = DOCTOR_SHEETS
REVIEW_SHEET = "Low (review)"
OTHER_PROFESSIONS_SHEET = "Other Professions"
# Why a sheet of the shortlist is never imported; any sheet not in DOCTOR_SHEETS is refused.
REFUSED_SHEETS = {
    OTHER_PROFESSIONS_SHEET: "students and non-doctors get no card and no profile (decided 20 Oct 2026)",
    REVIEW_SHEET: "its rows need a manual check first",
}
REPORT_PATH = OUTPUT_DIR / "seed_professionals_report.md"

# The registry columns read. Every other column of the shortlist (Address, Address2, Zip, Phone, Fax, Emails, Moved Since
# 2018, Old Address, Last Updated, aafia_score, aafia_tier, aafia_reasons) is deliberately left out.
REQUIRED_COLUMNS = ("NPI", "First Name", "Middle Name", "Last Name", "Credential", "Profession", "Taxonomy", "City", "State")
NOT_IMPORTED = (
    "Address", "Address2", "Zip", "Phone", "Fax", "Emails", "Moved Since 2018", "Old Address", "Last Updated",
    "aafia_score", "aafia_tier", "aafia_reasons",
)

# Professions whose display name takes the "Dr" title: the two doctor professions, the only ones imported.
DOCTOR_PROFESSIONS = frozenset({"physician", "dentist"})

# The taxonomy codes of each doctor profession, as NUCC v26.1 defines them. Physician: the grouping Allopathic &
# Osteopathic Physicians, which is exactly the codes starting 20. Dentist: the classification Dentist, which is the codes
# starting 1223 plus 125Q00000X (Oral Medicine). A unit test checks both rules against the NUCC grouping and classification
# of every shortlist code when the codes extract is available.
PHYSICIAN_CODE_PREFIX = "20"
DENTIST_CODE_PREFIX = "1223"
DENTIST_OTHER_CODES = frozenset({"125Q00000X"})

# USPS codes of the 50 states, DC and the inhabited territories plus the minor outlying islands. Military mail codes
# (AA, AE, AP) and foreign regions are not a practice state: state_code stays null and the report counts them. Must stay
# identical to private.is_us_state_code (supabase/migrations/20261020000300_l1_directory.sql); a unit test compares the
# two.
US_STATE_CODES = frozenset(
    """AL AK AZ AR CA CO CT DE FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA
    RI SC SD TN TX UT VT VA WA WV WI WY DC AS GU MP PR VI UM""".split()
)
# ISO 3166-1 alpha-2 country of a row whose State is one of the codes above (the territories included: the registry and
# the patient-facing card treat them as US practice).
US_COUNTRY = "US"

# The practice country of a doctor abroad (State not a US code), ISO 3166-1 alpha-2. Written from the State and City
# values actually found in the three doctor sheets of the October 2026 shortlist, keyed on the value as given in upper
# case with runs of spaces collapsed. A value missing from both tables gives null, never a guess, and the
# run prints it so the table can be extended when a new release brings one.
# A State that names one country on its own: a province, emirate, governorate, region or city of that country, or the
# country's name.
ABROAD_STATE_COUNTRY = {
    # Canada: provinces as the registry writes them ("ON" and "ONT." are Canada Post's and the common abbreviation).
    "ALBERTA": "CA", "MANITOBA": "CA", "ON": "CA", "ONT.": "CA", "ONTARIO": "CA", "QUEBEC": "CA",
    # Saudi Arabia: the Riyadh and Makkah regions, Najd (the central region) and Dhahran.
    "DHAHRAN": "SA", "MAKKAH": "SA", "MAKKAH REGION": "SA", "NEJD": "SA", "RIYADH": "SA", "SAUDI ARABIA": "SA",
    # United Arab Emirates: the emirates of Abu Dhabi and Dubai.
    "ABU DHABI": "AE", "DUBAI": "AE", "UNITED ARAB EMIRATES": "AE",
    # Kuwait: the Hawalli governorate.
    "HAWALLI": "KW", "KUWAIT": "KW",
    # Lebanon: the Beirut and North governorates.
    "BEIRUT": "LB", "NORTH LEBANON": "LB",
    # Egypt: the Alexandria and Ismailia governorates.
    "ALEXANDRIA": "EG", "ISMAILIA": "EG",
    "JORDAN": "JO",
    "PARIS": "FR",
    # Zimbabwe: Harare province.
    "HARARE": "ZW",
}
# A State that does not name one country on its own, settled by the City: (State, City) as given, upper case.
ABROAD_STATE_CITY_COUNTRY = {
    # "AE" is also a US military mail code; with the city Dubai (not APO, FPO or DPO) it is the ISO code of the UAE.
    ("AE", "DUBAI"): "AE",
    # "Central" and "Eastern Province" are region names in several countries; Riyadh and Dhahran are in Saudi Arabia.
    ("CENTRAL", "RIYADH"): "SA",
    ("EASTERN PROVINCE", "DHAHRAN"): "SA",
    # "Foreign" and "NA" (not applicable) say only that the practice is abroad.
    ("FOREIGN", "DUBAI"): "AE",
    ("NA", "DUBAI"): "AE",
    # KWI is Kuwait's airport code, not a region.
    ("KWI", "KUWAIT CITY"): "KW",
    # "SA" is also South Australia's code; Al-Khobar and Riyadh are in Saudi Arabia.
    ("SA", "AL-KHOBAR"): "SA",
    ("SA", "RIYADH"): "SA",
}
# US military mail "states" (Armed Forces Americas, Europe and the Middle East, Pacific): an APO, FPO or DPO address names
# no country, so the country stays null.
MILITARY_MAIL_STATES = frozenset({"AA", "AE", "AP"})

NPI_RE = re.compile(r"^[12][0-9]{9}$")
TAXONOMY_RE = re.compile(r"^[0-9]{3}[0-9A-Z]{6}X$")

# Column limits from the professionals table (supabase/migrations/20261020000300_l1_directory.sql).
NAME_MAX = 100
TEXT_MAX = 100

# The registry fields an update compares and writes, in table column order.
REGISTRY_FIELDS = (
    "first_name", "middle_name", "last_name", "name_en", "registry_credential", "registry_profession", "taxonomy_code",
    "city_en", "state_code", "practice_country",
)


@dataclasses.dataclass(frozen=True)
class RegistryRow:
    npi: str
    first_name: str
    middle_name: str | None
    last_name: str
    name_en: str
    registry_credential: str | None
    registry_profession: str | None
    taxonomy_code: str | None
    city_en: str | None
    state_code: str | None
    practice_country: str | None


@dataclasses.dataclass
class Reject:
    sheet: str
    row_number: int
    npi: str | None
    # A fixed text with no value from the row, so the report can count rejects by reason.
    reason: str
    # The row's own detail (the value refused, where a duplicate was first seen): stdout only.
    detail: str | None = None


@dataclasses.dataclass
class Shortlist:
    rows: list[RegistryRow]
    read_per_sheet: dict[str, int]
    accepted_per_sheet: dict[str, int]
    rejects: list[Reject]
    # Raw State values that are not a US state, DC or territory code, with their row counts (stdout only).
    non_us_states: collections.Counter
    blank_states: int
    # Well-formed rows of a doctor sheet that are not a doctor's (profession or taxonomy code): never imported.
    not_doctors: list[Reject] = dataclasses.field(default_factory=list)
    # (sheet, reason, profession, taxonomy code) -> rows: the report's aggregate of not_doctors.
    not_doctor_counts: collections.Counter = dataclasses.field(default_factory=collections.Counter)
    # Accepted rows by how their practice country was found (the keys of COUNTRY_BASES).
    country_bases: collections.Counter = dataclasses.field(default_factory=collections.Counter)
    # Raw (State, City) of accepted rows abroad whose country stays null and is not military mail (stdout only).
    abroad_unsettled: collections.Counter = dataclasses.field(default_factory=collections.Counter)

    def not_doctors_per_sheet(self, sheet: str) -> int:
        return sum(1 for r in self.not_doctors if r.sheet == sheet)

    def countries(self) -> collections.Counter:
        """Accepted rows per practice country (None: no country)."""
        return collections.Counter(r.practice_country for r in self.rows)


# How the practice country of an accepted row was found, in report order.
COUNTRY_BASES = (
    ("us", "US state, DC or territory code"),
    ("state", "abroad: the State names the country"),
    ("state_city", "abroad: the State and City name the country"),
    ("military", "abroad: US military mail code (AA, AE, AP), no country"),
    ("unknown", "abroad: the State and City name no one country, no country"),
    ("blank", "blank State, no country"),
)


# ---------------------------------------------------------------- row rules (unit-tested)


def npi_text(value: object) -> str | None:
    """The NPI as text. Spreadsheets may hand a number back, so whole floats lose their ".0"."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return clean(value)


def is_valid_npi(value: str | None) -> bool:
    return bool(value) and NPI_RE.fullmatch(value) is not None


def us_state_code(value: object) -> str | None:
    """The two-letter code when the registry State is a US state, DC or territory; otherwise None."""
    text = clean(value)
    if text is None:
        return None
    code = text.upper()
    return code if code in US_STATE_CODES else None


def _key(value: object) -> str | None:
    """A registry State or City as the country tables key it: upper case, runs of spaces collapsed."""
    text = clean(value)
    return re.sub(r"\s+", " ", text).upper() if text is not None else None


def country_basis(state_code: str | None, state: object = None, city: object = None) -> tuple[str | None, str]:
    """(ISO 3166-1 alpha-2 country, how it was found: a key of COUNTRY_BASES). 'US' for a US state, DC or territory code
    (as us_state_code returns it); for a row abroad, the country ABROAD_STATE_CITY_COUNTRY or ABROAD_STATE_COUNTRY names
    for the raw State (and City), else None: a military mail code, a blank State or a value no table lists."""
    if state_code in US_STATE_CODES:
        return US_COUNTRY, "us"
    state_key, city_key = _key(state), _key(city)
    if state_key is None:
        return None, "blank"
    by_city = ABROAD_STATE_CITY_COUNTRY.get((state_key, city_key))
    if by_city:
        return by_city, "state_city"
    if state_key in MILITARY_MAIL_STATES:
        return None, "military"
    by_state = ABROAD_STATE_COUNTRY.get(state_key)
    if by_state:
        return by_state, "state"
    return None, "unknown"


def practice_country(state_code: str | None, state: object = None, city: object = None) -> str | None:
    """The practice country: 'US' for a US state, DC or territory code; abroad, the country the registry's State (and
    City) name unambiguously, else None (see country_basis)."""
    return country_basis(state_code, state, city)[0]


def doctor_problem(profession: str | None, taxonomy: str | None) -> str | None:
    """None when the row is a doctor's: Profession Physician with a physician's taxonomy code, or Dentist with a
    dentist's code. Otherwise the reason the row is not imported (a fixed text: the report counts rows by it)."""
    kind = (profession or "").strip().casefold()
    if kind not in DOCTOR_PROFESSIONS:
        return "profession is not Physician or Dentist"
    code = (taxonomy or "").strip().upper()
    if not code:
        return "no taxonomy code"
    if kind == "physician" and not code.startswith(PHYSICIAN_CODE_PREFIX):
        return "Physician with a code outside NUCC Allopathic & Osteopathic Physicians (20...)"
    if kind == "dentist" and not (code.startswith(DENTIST_CODE_PREFIX) or code in DENTIST_OTHER_CODES):
        return "Dentist with a code outside the NUCC classification Dentist (1223..., 125Q00000X)"
    return None


def sheet_problem(sheet: str) -> str | None:
    """None for a doctor sheet; otherwise why the sheet is not imported."""
    if sheet in DOCTOR_SHEETS:
        return None
    reason = REFUSED_SHEETS.get(sheet, "it is not one of the doctor sheets")
    doctor_sheets = ", ".join(f'"{s}"' for s in DOCTOR_SHEETS)
    return f'"{sheet}" is not imported: {reason}. Only the doctor sheets are imported: {doctor_sheets}.'


def display_name(first_name: str, last_name: str, profession: str | None) -> str:
    """The card's English name: "Dr First Last" for physicians and dentists, "First Last" for every other profession."""
    name = f"{first_name} {last_name}"
    if profession is not None and profession.strip().casefold() in DOCTOR_PROFESSIONS:
        return f"Dr {name}"
    return name


def parse_row(values: dict[str, object]) -> tuple[RegistryRow | None, str | None]:
    """One shortlist row to a RegistryRow, or (None, reason) when the row cannot be imported. The reason never quotes a
    value of the row (the report counts rejects by reason)."""
    npi = npi_text(values.get("NPI"))
    if not is_valid_npi(npi):
        return None, "NPI is not 10 digits starting with 1 or 2"
    first_name = clean(values.get("First Name"))
    last_name = clean(values.get("Last Name"))
    if first_name is None or last_name is None:
        return None, "first or last name is blank"
    middle_name = clean(values.get("Middle Name"))
    credential = clean(values.get("Credential"))
    profession = clean(values.get("Profession"))
    city = clean(values.get("City"))
    taxonomy = clean(values.get("Taxonomy"))
    if taxonomy is not None:
        taxonomy = taxonomy.upper()
        if TAXONOMY_RE.fullmatch(taxonomy) is None:
            return None, "taxonomy code is not a NUCC code"
    for label, text, limit in (
        ("first name", first_name, NAME_MAX),
        ("middle name", middle_name, NAME_MAX),
        ("last name", last_name, NAME_MAX),
        ("credential", credential, TEXT_MAX),
        ("profession", profession, TEXT_MAX),
        ("city", city, TEXT_MAX),
    ):
        if text is not None and len(text) > limit:
            return None, f"{label} is longer than {limit} characters"
    state_code = us_state_code(values.get("State"))
    row = RegistryRow(
        npi=npi,
        first_name=first_name,
        middle_name=middle_name,
        last_name=last_name,
        name_en=display_name(first_name, last_name, profession),
        registry_credential=credential,
        registry_profession=profession,
        taxonomy_code=taxonomy,
        city_en=city,
        state_code=state_code,
        practice_country=practice_country(state_code, values.get("State"), city),
    )
    return row, None


# ---------------------------------------------------------------- reading the workbook


def read_shortlist(path: pathlib.Path, sheets: Sequence[str]) -> Shortlist:
    """Reads the named doctor sheets in order. A row that is not a doctor's is counted, never imported; an NPI seen in an
    earlier sheet is rejected as a duplicate in a later one. Refuses any sheet that is not a doctor sheet."""
    import openpyxl  # imported here so --help runs without it

    refused = [problem for problem in map(sheet_problem, sheets) if problem]
    if refused:
        raise SystemExit("\n".join(refused))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [name for name in sheets if name not in workbook.sheetnames]
        if missing:
            raise SystemExit(f"Sheet(s) not in {path.name}: {', '.join(missing)} (has: {', '.join(workbook.sheetnames)})")
        result = Shortlist([], {}, {}, [], collections.Counter(), 0)
        seen: dict[str, tuple[str, int]] = {}
        for sheet in sheets:
            rows = workbook[sheet].iter_rows(values_only=True)
            header = [clean(h) for h in next(rows, ())]
            absent = [c for c in REQUIRED_COLUMNS if c not in header]
            if absent:
                raise SystemExit(f'Sheet "{sheet}" lacks column(s): {", ".join(absent)}')
            read = accepted = 0
            for row_number, cells in enumerate(rows, start=2):
                if cells is None or all(clean(c) is None for c in cells):
                    continue
                read += 1
                values = dict(zip(header, cells))
                parsed, reason = parse_row(values)
                if parsed is None:
                    result.rejects.append(Reject(sheet, row_number, npi_text(values.get("NPI")), reason or "invalid"))
                    continue
                not_doctor = doctor_problem(parsed.registry_profession, parsed.taxonomy_code)
                if not_doctor:
                    profession, code = parsed.registry_profession or "(blank)", parsed.taxonomy_code or "(blank)"
                    result.not_doctors.append(Reject(sheet, row_number, parsed.npi, not_doctor, f"{profession} {code}"))
                    result.not_doctor_counts[(sheet, not_doctor, profession, code)] += 1
                    continue
                if parsed.npi in seen:
                    first_sheet, first_row = seen[parsed.npi]
                    result.rejects.append(Reject(sheet, row_number, parsed.npi, "duplicate NPI (first sheet wins)",
                                                 f'first seen in "{first_sheet}" row {first_row}'))
                    continue
                seen[parsed.npi] = (sheet, row_number)
                raw_state, raw_city = clean(values.get("State")), clean(values.get("City"))
                _, basis = country_basis(parsed.state_code, raw_state, raw_city)
                if raw_state is None:
                    result.blank_states += 1
                elif parsed.state_code is None:
                    result.non_us_states[raw_state] += 1
                    if basis == "unknown":
                        result.abroad_unsettled[(raw_state, raw_city or "")] += 1
                result.country_bases[basis] += 1
                result.rows.append(parsed)
                accepted += 1
            result.read_per_sheet[sheet] = read
            result.accepted_per_sheet[sheet] = accepted
        return result
    finally:
        workbook.close()


# ---------------------------------------------------------------- database


def _copy_stage(cur, rows: Sequence[RegistryRow], batch_size: int) -> None:
    columns = ("npi",) + REGISTRY_FIELDS
    sql = f"copy registry_stage ({', '.join(columns)}) from stdin with (format csv)"
    for start in range(0, len(rows), batch_size):
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        for row in rows[start : start + batch_size]:
            # None becomes an unquoted empty field, which COPY reads as NULL (no field here is ever an empty string).
            writer.writerow([getattr(row, c) for c in columns])
        buffer.seek(0)
        cur.copy_expert(sql, buffer)


def import_rows(conn, rows: Sequence[RegistryRow], registry_as_of: dt.date, batch_size: int = 5000) -> dict:
    """Stages the rows, upserts professionals and their primary specialty links. Does not commit."""
    stats: dict = {}
    with conn.cursor() as cur:
        cur.execute("drop table if exists pg_temp.registry_stage, pg_temp.registry_scope")
        cur.execute(
            """
            create temp table registry_stage (
              npi text primary key,
              first_name text not null,
              middle_name text,
              last_name text not null,
              name_en text not null,
              registry_credential text,
              registry_profession text,
              taxonomy_code text,
              city_en text,
              state_code text,
              practice_country text,
              specialty_id uuid
            ) on commit drop
            """
        )
        _copy_stage(cur, rows, batch_size)

        # The group each code maps to. A code maps to one group only (trigger on specialties); unmapped stays null.
        cur.execute(
            """
            update registry_stage s set specialty_id = m.id
              from (select id, unnest(taxonomy_codes) as code from public.specialties) m
             where m.code = s.taxonomy_code
            """
        )
        cur.execute(
            """
            select count(*),
                   count(*) filter (where specialty_id is not null),
                   (select count(*) from public.specialties),
                   (select count(*) from public.specialties where cardinality(taxonomy_codes) > 0)
              from registry_stage
            """
        )
        staged, mapped, specialties, specialties_with_codes = cur.fetchone()
        stats.update(staged=staged, mapped=mapped, unmapped=staged - mapped, specialties=specialties,
                     specialties_with_codes=specialties_with_codes)
        cur.execute(
            """
            select coalesce(taxonomy_code, '(none)'), min(registry_profession), count(*)
              from registry_stage where specialty_id is null
             group by 1 order by 3 desc, 1 limit 15
            """
        )
        stats["top_unmapped"] = cur.fetchall()

        # Rows of this shortlist that already exist, split into the ones the import may refresh and the ones it leaves.
        cur.execute(
            """
            select p.status::text, p.source::text, count(*)
              from registry_stage s join public.professionals p on p.npi = s.npi
             group by 1, 2 order by 1, 2
            """
        )
        existing = cur.fetchall()
        stats["existing_open"] = sum(n for status, source, n in existing if status == "unclaimed" and source == "nppes")
        stats["skipped"] = [(status, source, n) for status, source, n in existing
                            if not (status == "unclaimed" and source == "nppes")]

        set_list = ",\n                ".join(f"{c} = excluded.{c}" for c in REGISTRY_FIELDS + ("specialty_id",))
        # The compared tuple: the registry facts, the specialty they map to and the release date, so a newer release with
        # the same facts moves registry_as_of while a re-run of the same file changes nothing.
        compared = REGISTRY_FIELDS + ("specialty_id", "registry_as_of")
        old = ", ".join(f"p.{c}" for c in compared)
        new = ", ".join(f"excluded.{c}" for c in compared)
        cur.execute(
            f"""
            with upserted as (
              insert into public.professionals as p (
                npi, status, source, {', '.join(REGISTRY_FIELDS)}, specialty_id, registry_source, registry_as_of
              )
              select npi, 'unclaimed', 'nppes', {', '.join(REGISTRY_FIELDS)}, specialty_id, 'nppes', %(as_of)s
                from registry_stage
               order by npi
              on conflict (npi) do update set
                {set_list},
                registry_source = 'nppes',
                registry_as_of = excluded.registry_as_of
              where p.status = 'unclaimed' and p.source = 'nppes'
                and ({old}) is distinct from ({new})
              returning (xmax = 0) as inserted
            )
            select count(*) filter (where inserted), count(*) filter (where not inserted) from upserted
            """,
            {"as_of": registry_as_of},
        )
        stats["inserted"], stats["updated"] = cur.fetchone()
        stats["unchanged"] = stats["existing_open"] - stats["updated"]

        # Primary specialty links of the unclaimed registry rows in this shortlist follow specialty_id. Links of claimed
        # rows, and any non-primary link, are left alone.
        cur.execute(
            """
            create temp table registry_scope on commit drop as
            select p.id, p.specialty_id, p.taxonomy_code
              from public.professionals p join registry_stage s on s.npi = p.npi
             where p.status = 'unclaimed' and p.source = 'nppes'
            """
        )
        cur.execute(
            """
            delete from public.professional_specialties ps
             using registry_scope r
             where ps.professional_id = r.id and ps.is_primary and ps.specialty_id is distinct from r.specialty_id
            """
        )
        stats["links_deleted"] = cur.rowcount
        cur.execute(
            """
            with linked as (
              insert into public.professional_specialties as ps (professional_id, specialty_id, taxonomy_code, is_primary)
              select id, specialty_id, taxonomy_code, true from registry_scope where specialty_id is not null
              on conflict (professional_id, specialty_id) do update
                 set is_primary = true, taxonomy_code = excluded.taxonomy_code
               where (ps.is_primary, ps.taxonomy_code) is distinct from (true, excluded.taxonomy_code)
              returning (xmax = 0) as inserted
            )
            select count(*) filter (where inserted), count(*) filter (where not inserted) from linked
            """
        )
        stats["links_inserted"], stats["links_updated"] = cur.fetchone()
        cur.execute("drop table registry_scope")

        cur.execute(
            """
            select count(*),
                   count(*) filter (where p.practice_country = %(us)s),
                   count(*) filter (where p.practice_country is null),
                   count(*) filter (where p.specialty_id is not null),
                   count(*) filter (where exists (select 1 from public.professional_specialties ps
                                                   where ps.professional_id = p.id and ps.is_primary))
              from public.professionals p join registry_stage s on s.npi = p.npi
             where p.source = 'nppes'
            """,
            {"us": US_COUNTRY},
        )
        (stats["final_nppes_rows"], stats["final_country_us"], stats["final_country_null"],
         stats["final_with_specialty"], stats["final_with_primary_link"]) = cur.fetchone()
        cur.execute(
            """
            select p.practice_country, count(*)
              from public.professionals p join registry_stage s on s.npi = p.npi
             where p.source = 'nppes'
             group by 1
            """
        )
        stats["final_countries"] = dict(cur.fetchall())
        cur.execute("drop table registry_stage")
    return stats


# ---------------------------------------------------------------- report


def describe_database(url: str | None) -> str:
    """host:port/dbname of a connection string in either libpq form (URL or key=value), never the user or the password;
    "(unparsed)" when it cannot be read, never the raw string (the reports that print this are committed)."""
    if not url:
        return "(none)"
    try:
        from psycopg2.extensions import parse_dsn  # imported here so the unit tests run without the driver installed

        parts = parse_dsn(url)
    except Exception:  # any failure (no driver, a malformed string) gives the same safe answer
        return "(unparsed)"
    host = parts.get("host") or parts.get("hostaddr") or "?"
    port = f":{parts['port']}" if parts.get("port") else ""
    dbname = f"/{parts['dbname']}" if parts.get("dbname") else ""
    return f"{host}{port}{dbname}"


def file_digest(path: pathlib.Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def previous_run_log(path: pathlib.Path, columns: int | None = None) -> list[str]:
    """The run log rows of the last report, so the file keeps the history of real runs (first and second run). With
    `columns`, rows written with a different set of columns (an older version of the script) are dropped."""
    if not path.exists():
        return []
    lines = path.read_text(encoding="utf-8").splitlines()
    try:
        start = lines.index("## Run log")
    except ValueError:
        return []
    rows = [line for line in lines[start + 1 :] if line.startswith("| 20")]
    if columns is not None:
        rows = [line for line in rows if line.count("|") - 1 == columns]
    return rows


def report_lines(args, shortlist: Shortlist, stats: dict, started: dt.datetime, seconds: float,
                 run_log: list[str]) -> list[str]:
    total_read = sum(shortlist.read_per_sheet.values())
    total_accepted = len(shortlist.rows)
    total_not_doctors = len(shortlist.not_doctors)
    nppes_equal = stats["final_nppes_rows"] == total_accepted
    refused = ", ".join(f'"{s}"' for s in REFUSED_SHEETS)
    lines = [
        "# Professional import report (task 1.5)",
        "",
        f"Generated by `scripts/seed_professionals.py` at {started.strftime('%Y-%m-%d %H:%M UTC')}"
        f" ({'dry run, rolled back' if args.dry_run else 'committed'}, {seconds:.1f} s). Public NPPES registry data,"
        " loaded into the local database only.",
        "",
        "Only doctors get a card and a profile (decided 20 Oct 2026): the doctor sheets are the import, and within them a"
        " row is imported only when its Profession is Physician with a physician's taxonomy code (NUCC Allopathic &"
        " Osteopathic Physicians, 20...) or Dentist with a dentist's code (NUCC classification Dentist, 1223... and"
        " 125Q00000X).",
        "",
        "## Inputs",
        "",
        f"- Shortlist: `{args.shortlist.name}` ({args.shortlist.stat().st_size:,} bytes, sha256"
        f" `{file_digest(args.shortlist)[:16]}...`)",
        "- Sheets: " + ", ".join(f'"{s}"' for s in args.sheets) + f" (the doctor sheets; never {refused})",
        f"- Registry as of: {args.registry_as_of.isoformat()} (NPPES release date, stored as `registry_as_of`)",
        f"- Database: `{describe_database(args.database_url)}`",
        "",
        "## Rows read",
        "",
        *markdown_table(
            ["Sheet", "Rows read", "Imported", "Not a doctor's row", "Rejected"],
            [
                (s, f"{shortlist.read_per_sheet[s]:,}", f"{shortlist.accepted_per_sheet[s]:,}",
                 f"{shortlist.not_doctors_per_sheet(s):,}",
                 f"{shortlist.read_per_sheet[s] - shortlist.accepted_per_sheet[s] - shortlist.not_doctors_per_sheet(s):,}")
                for s in args.sheets
            ]
            + [("Total", f"{total_read:,}", f"{total_accepted:,}", f"{total_not_doctors:,}",
                f"{total_read - total_accepted - total_not_doctors:,}")],
        ),
        "",
    ]
    if shortlist.not_doctor_counts:
        # Aggregates only: the rows themselves (NPI, row number) go to stdout, never into the committed report.
        lines += ["Rows of the doctor sheets that are not a doctor's, never imported (the rows are printed to stdout"
                  " only):", ""]
        lines += markdown_table(
            ["Sheet", "Reason", "Profession", "Taxonomy code", "Rows"],
            [(sheet, reason, profession, code, f"{n:,}")
             for (sheet, reason, profession, code), n in sorted(shortlist.not_doctor_counts.items())],
        )
        lines.append("")
    else:
        lines += ["Every well-formed row of the doctor sheets is a doctor's (Profession and taxonomy code agree).", ""]
    if shortlist.rejects:
        # Counts only: the rejected rows themselves (NPI, row number) go to stdout, never into the committed report.
        by_reason = collections.Counter((r.sheet, r.reason) for r in shortlist.rejects)
        lines += ["Rejected rows by reason (the rows themselves are printed to stdout only):", ""]
        lines += markdown_table(
            ["Sheet", "Reason", "Rows"],
            [(sheet, reason, f"{n:,}") for (sheet, reason), n in sorted(by_reason.items())],
        )
        lines.append("")
    else:
        lines += ["No row was rejected (every NPI matches `^[12][0-9]{9}$`, names present, taxonomy codes NUCC-shaped).", ""]

    skipped_total = sum(n for _, _, n in stats["skipped"])
    lines += [
        "## Database changes",
        "",
        *markdown_table(
            ["Inserted", "Updated", "Unchanged", "Skipped (claimed or not from nppes)"],
            [(f"{stats['inserted']:,}", f"{stats['updated']:,}", f"{stats['unchanged']:,}", f"{skipped_total:,}")],
        ),
        "",
    ]
    if stats["skipped"]:
        lines += markdown_table(["Status", "Source", "Rows left untouched"], stats["skipped"]) + [""]
    lines += [
        f"Primary specialty links (`professional_specialties`, unclaimed nppes rows only): {stats['links_inserted']:,}"
        f" inserted, {stats['links_updated']:,} updated, {stats['links_deleted']:,} deleted.",
        "",
        "An existing row changes only while it is unclaimed and from nppes and an imported value or the registry date"
        " differs (IS DISTINCT FROM; a newer release with the same facts only moves `registry_as_of`); its slug is never"
        " rewritten. A run with the same file changes zero rows.",
        "",
        "## Specialty mapping",
        "",
        f"- Specialty groups in the database: {stats['specialties']:,}, of which {stats['specialties_with_codes']:,}"
        " carry taxonomy codes (`map_taxonomy.py`).",
        f"- Rows with a mapped specialty: {stats['mapped']:,}; unmapped (specialty_id null, never defaulted):"
        f" {stats['unmapped']:,}.",
        "",
    ]
    if stats["top_unmapped"]:
        lines += ["Largest unmapped taxonomy codes in this shortlist:", ""]
        lines += markdown_table(["Taxonomy code", "Profession", "Rows"], [(c, p or "", f"{n:,}") for c, p, n in
                                                                          stats["top_unmapped"]])
        lines.append("")
    countries = shortlist.countries()
    country_us = countries.get(US_COUNTRY, 0)
    country_null = countries.get(None, 0)
    final_countries = stats.get("final_countries", {})
    lines += [
        "## Practice state and country",
        "",
        "A US state, DC or territory code is `state_code` with practice_country `US` (the code list"
        " `private.is_us_state_code` checks: 50 states, DC, AS GU MP PR VI UM). Abroad, `state_code` is null and the"
        " country is the one the registry's State, or State and City, names unambiguously (the tables in"
        " `seed_professionals.py`); null where they name none, as for the US military mail codes AA, AE and AP.",
        "",
        *markdown_table(
            ["How the country was found", "Rows"],
            [(label, f"{shortlist.country_bases.get(key, 0):,}") for key, label in COUNTRY_BASES],
        ),
        "",
        "Rows per practice country (ISO 3166-1 alpha-2):",
        "",
        *markdown_table(
            ["Country", "Rows in the file", "Rows in the database after the run"],
            [(country or "(none)", f"{n:,}", f"{final_countries.get(country, 0):,}")
             for country, n in sorted(countries.items(), key=lambda item: (-item[1], item[0] or "~"))],
        ),
        "",
        "## Check",
        "",
        f"Doctor rows imported (`source = nppes` rows whose NPI is in this import, after the run):"
        f" {stats['final_nppes_rows']:,}; doctor rows of the shortlist that pass the guard: {total_accepted:,};"
        f" equal: {'yes' if nppes_equal else 'NO'}.",
        "",
        *markdown_table(
            ["After the run (nppes rows of this import)", "Rows", "Expected from the file"],
            [
                (f"practice_country = {US_COUNTRY}", f"{stats['final_country_us']:,}", f"{country_us:,}"),
                ("practice_country abroad (another country)",
                 f"{stats['final_nppes_rows'] - stats['final_country_us'] - stats['final_country_null']:,}",
                 f"{total_accepted - country_us - country_null:,}"),
                ("practice_country null", f"{stats['final_country_null']:,}", f"{country_null:,}"),
                ("specialty_id set", f"{stats['final_with_specialty']:,}", f"{stats['mapped']:,} (mapped)"),
                ("primary professional_specialties link", f"{stats['final_with_primary_link']:,}",
                 f"{stats['mapped']:,} (mapped)"),
            ],
        ),
        "",
        "Claimed or self-listed rows of the shortlist (none on a fresh database) keep whatever they hold, so on a later"
        " run the first column may differ from the file by the skipped rows. Rows an earlier import wrote outside the"
        " doctor scope are not deleted by this script.",
        "",
        "## Not imported",
        "",
        "Columns left out on purpose: " + ", ".join(f"`{c}`" for c in NOT_IMPORTED) + ". An unclaimed page carries no"
        " contact details (ADR-5); outreach emails are a separate decision; the scores are internal to the shortlist.",
        "",
        "## Run log",
        "",
        "| " + " | ".join(RUN_LOG_COLUMNS) + " |",
        "|" + "|".join(" --- " for _ in RUN_LOG_COLUMNS) + "|",
        *run_log,
    ]
    return lines


RUN_LOG_COLUMNS = ("Run at (UTC)", "Registry as of", "Doctor rows imported", "Not doctors' rows", "Inserted", "Updated",
                   "Unchanged", "Skipped", "Mapped", "Links changed")


def summary_line(args, shortlist: Shortlist, stats: dict, started: dt.datetime) -> str:
    links = stats["links_inserted"] + stats["links_updated"] + stats["links_deleted"]
    return (
        f"| {started.strftime('%Y-%m-%d %H:%M')} | {args.registry_as_of.isoformat()} | {len(shortlist.rows):,}"
        f" | {len(shortlist.not_doctors):,} | {stats['inserted']:,} | {stats['updated']:,} | {stats['unchanged']:,}"
        f" | {sum(n for _, _, n in stats['skipped']):,} | {stats['mapped']:,} | {links:,} |"
    )


# ---------------------------------------------------------------- command line


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--shortlist", required=True, type=pathlib.Path, help="aafia_registry_shortlist_YYYY-MM.xlsx")
    parser.add_argument("--sheets", nargs="+", default=list(DEFAULT_SHEETS),
                        help=f"Doctor sheets to import, in order (default: all three: {', '.join(DEFAULT_SHEETS)})."
                        f" '{OTHER_PROFESSIONS_SHEET}', '{REVIEW_SHEET}' and any other sheet are refused.")
    parser.add_argument("--registry-as-of", required=True, type=dt.date.fromisoformat,
                        help="NPPES release date of the shortlist, YYYY-MM-DD (stored as registry_as_of)")
    add_database_argument(parser)
    parser.add_argument("--dry-run", action="store_true", help="Run everything, report, then roll back")
    parser.add_argument("--batch-size", type=int, default=5000, help="Rows per COPY into the staging table")
    args = parser.parse_args(argv)
    refused = [problem for problem in map(sheet_problem, args.sheets) if problem]
    if refused:
        parser.error(" ".join(refused))
    if len(set(args.sheets)) != len(args.sheets):
        parser.error("--sheets names a sheet twice")
    if args.batch_size < 1:
        parser.error("--batch-size must be at least 1")
    if not args.shortlist.is_file():
        parser.error(f"no such file: {args.shortlist}")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    started = utc_now()
    clock = time.monotonic()
    shortlist = read_shortlist(args.shortlist, args.sheets)
    print(f"Read {sum(shortlist.read_per_sheet.values()):,} rows from {len(args.sheets)} doctor sheets"
          f" ({len(shortlist.rows):,} doctor rows to import, {len(shortlist.not_doctors):,} not a doctor's,"
          f" {len(shortlist.rejects):,} rejected).")
    for r in shortlist.not_doctors[:50]:
        print(f'  not imported: "{r.sheet}" row {r.row_number}, NPI {r.npi or "-"}: {r.reason} ({r.detail})')
    for r in shortlist.rejects[:50]:
        print(f'  rejected: "{r.sheet}" row {r.row_number}, NPI {r.npi or "-"}: {r.reason}'
              + (f" ({r.detail})" if r.detail else ""))
    for (state, city), n in sorted(shortlist.abroad_unsettled.items()):
        print(f'  abroad, no country: State "{state}", City "{city}": {n} row(s) (add it to the country tables if it'
              " names one country)")
    conn = connect(args.database_url)
    try:
        stats = import_rows(conn, shortlist.rows, args.registry_as_of, args.batch_size)
        if args.dry_run:
            conn.rollback()
        else:
            conn.commit()
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()
    seconds = time.monotonic() - clock
    run_log = previous_run_log(REPORT_PATH, len(RUN_LOG_COLUMNS))
    if not args.dry_run:
        run_log = (run_log + [summary_line(args, shortlist, stats, started)])[-20:]
    lines = report_lines(args, shortlist, stats, started, seconds, run_log)
    print("\n".join(lines))
    if args.dry_run:
        print("\nDry run: nothing committed, report file not written.")
    else:
        print(f"\nReport written to {write_report(REPORT_PATH, lines)}")
    return 0 if stats["final_nppes_rows"] == len(shortlist.rows) else 1


if __name__ == "__main__":
    sys.exit(main())
