"""Maps NPPES (NUCC) taxonomy codes to the 104 specialty groups and writes specialties.taxonomy_codes (task 1.5, L1).

The mapping itself is authored in scripts/taxonomy_map.py. This script counts the registry rows per code in its input,
writes output/taxonomy_exceptions.csv with every code it meets that TAXONOMY_MAP does not map (with UNMAPPED's reason,
or "not yet in taxonomy_map.py" for a code in neither dict: an unmapped code is listed, never defaulted to a group),
then writes each group's codes into the database in one transaction.

Input is either the shortlist workbook or a CSV with a code column and an optional count column. From the workbook it
reads the rows the registry import takes: the three doctor sheets ("Physicians High", "Physicians Medium", "Dentists",
28,587 rows; only doctors get a card, decided 20 Oct 2026), less the rows seed_professionals.py does not import because
their Profession and taxonomy code are not a doctor's (counted in the report, not in the mapping). The non-doctor codes
of the full shortlist keep their UNMAPPED reasons in taxonomy_map.py; --codes-csv with the 589-code extract lists them.
--nucc adds the NUCC grouping, classification and specialization to the exceptions file.

Database: every code in TAXONOMY_MAP goes to its group when the group exists (run seed_specialties.py first). Each
changed group is cleared before any is set, so a code can move between groups without tripping the one-group-per-code
trigger. Only rows whose codes actually change are written, so a second run reports 0 changed. Groups outside the 104
are left alone; if one of them holds a code TAXONOMY_MAP gives to a group, the run stops before writing.

    python scripts/map_taxonomy.py --shortlist PATH [--nucc PATH] [--dry-run]
    python scripts/map_taxonomy.py --codes-csv PATH [--nucc PATH] [--dry-run]

Writes output/taxonomy_exceptions.csv (built from the input file alone, so a dry run writes it too) and, on a committed
run, its section of output/specialties_report.md; a dry run prints that section instead, so the committed report always
describes the last committed run.
"""

from __future__ import annotations

import argparse
import csv
import pathlib
import re
import sys
from collections import Counter
from collections.abc import Iterable, Sequence
from typing import Any

# Importable under `python -I` too (isolated mode leaves the script's folder off sys.path).
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from seed_common import OUTPUT_DIR, add_database_argument, connect, markdown_table, utc_now  # noqa: E402
import seed_specialties  # noqa: E402
from seed_professionals import DOCTOR_SHEETS, doctor_problem  # noqa: E402
from seed_specialties import specialty_rows, write_report_section  # noqa: E402
from taxonomy_map import TAXONOMY_MAP, UNMAPPED  # noqa: E402

EXCEPTIONS_PATH = OUTPUT_DIR / "taxonomy_exceptions.csv"
# The sheets the registry import takes (seed_professionals.py): the doctor sheets.
SHORTLIST_SHEETS = DOCTOR_SHEETS
CODE_PATTERN = re.compile(r"^[0-9]{3}[0-9A-Z]{6}X$")
NOT_IN_MAP = "not yet in taxonomy_map.py"
BLANK = "(blank)"
EXCEPTION_COLUMNS = ("code", "rows", "profession", "grouping", "classification", "specialization", "reason")


class InputError(Exception):
    pass


def check_mapping(slugs: Iterable[str]) -> None:
    """Refuses to run on a broken mapping file (the unit tests check the same things)."""
    known = set(slugs)
    problems = [
        f"{code}: slug {slug!r} is not one of the 104 groups"
        for code, slug in TAXONOMY_MAP.items()
        if slug not in known
    ]
    problems += [f"{code}: in both TAXONOMY_MAP and UNMAPPED" for code in sorted(set(TAXONOMY_MAP) & set(UNMAPPED))]
    problems += [
        f"{code}: not a NUCC code" for code in sorted(set(TAXONOMY_MAP) | set(UNMAPPED)) if not CODE_PATTERN.match(code)
    ]
    if problems:
        raise ValueError("scripts/taxonomy_map.py is inconsistent:\n  " + "\n  ".join(problems))


def normalise_code(raw: object) -> str:
    text = str(raw).strip().upper() if raw is not None else ""
    return text or BLANK


# ---------------------------------------------------------------- inputs


def read_shortlist(path: pathlib.Path) -> tuple[Counter, dict[str, Counter], Counter]:
    """Rows per code and the shortlist's Profession values per code, from the doctor sheets, counting only the rows
    seed_professionals.py imports; the third value counts the other rows by (profession, code)."""
    import openpyxl  # imported here so the unit tests and --codes-csv runs do not need it

    counts: Counter = Counter()
    professions: dict[str, Counter] = {}
    not_doctors: Counter = Counter()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [name for name in SHORTLIST_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise InputError(f"{path.name} has no sheet {', '.join(repr(m) for m in missing)}")
        for name in SHORTLIST_SHEETS:
            rows = workbook[name].iter_rows(values_only=True)
            header = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
            if "Taxonomy" not in header:
                raise InputError(f"sheet {name!r} has no Taxonomy column")
            code_at = header.index("Taxonomy")
            profession_at = header.index("Profession") if "Profession" in header else None
            for row in rows:
                if row is None or all(cell is None for cell in row):
                    continue
                code = normalise_code(row[code_at] if code_at < len(row) else None)
                cell = row[profession_at] if profession_at is not None and profession_at < len(row) else None
                profession = str(cell).strip() if cell is not None and str(cell).strip() else None
                if doctor_problem(profession, None if code == BLANK else code):
                    not_doctors[(profession or "(blank)", code)] += 1
                    continue
                counts[code] += 1
                if profession:
                    professions.setdefault(code, Counter())[profession] += 1
    finally:
        workbook.close()
    return counts, professions, not_doctors


def read_codes_csv(path: pathlib.Path) -> tuple[Counter, dict[str, Counter], dict[str, dict[str, str]]]:
    """A CSV with a code (or Taxonomy) column, an optional count (or rows) column and, optionally, profession and the
    NUCC grouping, classification and specialization (the shape of the shortlist codes extract)."""
    counts: Counter = Counter()
    professions: dict[str, Counter] = {}
    described: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {name.strip().lower(): name for name in reader.fieldnames or []}
        code_field = fields.get("code") or fields.get("taxonomy")
        if not code_field:
            raise InputError(f"{path.name} has no code or Taxonomy column")
        count_field = fields.get("count") or fields.get("rows")
        for row in reader:
            code = normalise_code(row.get(code_field))
            try:
                rows = int(row[count_field]) if count_field and row.get(count_field) not in (None, "") else 1
            except ValueError:
                raise InputError(f"{path.name}: count {row[count_field]!r} for {code} is not a whole number") from None
            counts[code] += rows
            if fields.get("profession") and row.get(fields["profession"]):
                professions.setdefault(code, Counter())[row[fields["profession"]].strip()] += rows
            extra = {
                key: (row.get(fields[key]) or "").strip()
                for key in ("grouping", "classification", "specialization")
                if key in fields
            }
            if any(extra.values()):
                described[code] = extra
    return counts, professions, described


def read_nucc(path: pathlib.Path) -> dict[str, dict[str, str]]:
    """The NUCC code set CSV (Code, Grouping, Classification, Specialization, ...)."""
    out: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or "Code" not in reader.fieldnames:
            raise InputError(f"{path.name} has no Code column; expected the NUCC taxonomy CSV")
        for row in reader:
            out[row["Code"].strip().upper()] = {
                "grouping": (row.get("Grouping") or "").strip(),
                "classification": (row.get("Classification") or "").strip(),
                "specialization": (row.get("Specialization") or "").strip(),
            }
    return out


# ---------------------------------------------------------------- exceptions


def exception_reason(code: str) -> str:
    if code == BLANK:
        return "no taxonomy code on the row"
    if not CODE_PATTERN.match(code):
        return "not a NUCC taxonomy code"
    return UNMAPPED.get(code, NOT_IN_MAP)


def build_exceptions(
    counts: Counter,
    professions: dict[str, Counter],
    descriptions: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    """One row per input code that TAXONOMY_MAP does not map, most rows first."""
    rows = []
    for code, total in counts.items():
        if code in TAXONOMY_MAP:
            continue
        info = descriptions.get(code, {})
        profession = professions.get(code)
        rows.append(
            {
                "code": code,
                "rows": total,
                "profession": profession.most_common(1)[0][0] if profession else "",
                "grouping": info.get("grouping", ""),
                "classification": info.get("classification", ""),
                "specialization": info.get("specialization", ""),
                "reason": exception_reason(code),
            }
        )
    rows.sort(key=lambda row: (-row["rows"], row["code"]))
    return rows


def write_exceptions(rows: Sequence[dict[str, Any]], path: pathlib.Path = EXCEPTIONS_PATH) -> pathlib.Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXCEPTION_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return path


# ---------------------------------------------------------------- database


def desired_codes(slugs: Iterable[str]) -> dict[str, list[str]]:
    """Each group's codes from TAXONOMY_MAP, sorted (the stored order, so a re-run compares equal)."""
    out: dict[str, list[str]] = {slug: [] for slug in slugs}
    for code, slug in TAXONOMY_MAP.items():
        out.setdefault(slug, []).append(code)
    return {slug: sorted(codes) for slug, codes in out.items()}


def apply_to_database(conn, slugs: Sequence[str]) -> dict[str, Any]:
    """Writes specialties.taxonomy_codes for the 104 groups. The caller commits or rolls back."""
    wanted = desired_codes(slugs)
    with conn.cursor() as cur:
        cur.execute("select id, slug, taxonomy_codes from public.specialties order by position, slug for update")
        present = {slug: (row_id, list(codes or [])) for row_id, slug, codes in cur.fetchall()}

        ours = set(slugs)
        missing_groups = [slug for slug in slugs if slug not in present]
        skipped_codes = sorted(code for slug in missing_groups for code in wanted[slug])
        # Codes we are about to give a group but which a group outside the 104 already holds: stop, change nothing.
        writable = {code for slug in slugs if slug in present for code in wanted[slug]}
        conflicts = sorted(
            (code, slug)
            for slug, (_, codes) in present.items()
            if slug not in ours
            for code in codes
            if code in writable
        )
        if conflicts:
            listed = ", ".join(f"{code} (held by {slug})" for code, slug in conflicts)
            raise RuntimeError(f"codes held by groups outside the 104, nothing written: {listed}")

        changes = [
            (present[slug][0], slug, present[slug][1], wanted[slug])
            for slug in slugs
            if slug in present and present[slug][1] != wanted[slug]
        ]
        if changes:
            cur.execute(
                "update public.specialties set taxonomy_codes = '{}'"
                " where id = any(%s::uuid[]) and taxonomy_codes <> '{}'",
                ([row_id for row_id, *_ in changes],),
            )
            for row_id, _, _, codes in changes:
                if codes:
                    cur.execute(
                        "update public.specialties set taxonomy_codes = %s::text[] where id = %s", (codes, row_id)
                    )

        cur.execute(
            "select count(*) from public.specialties where cardinality(taxonomy_codes) > 0 and slug = any(%s)",
            (list(slugs),),
        )
        groups_with_codes = cur.fetchone()[0]

    return {
        "changed": [
            {"slug": slug, "added": sorted(set(new) - set(old)), "removed": sorted(set(old) - set(new))}
            for _, slug, old, new in changes
        ],
        "missing_groups": missing_groups,
        "skipped_codes": skipped_codes,
        "groups_with_codes": groups_with_codes,
        "codes_written": sum(len(wanted[slug]) for slug in slugs if slug in present),
    }


# ---------------------------------------------------------------- summary and report


def summarise(counts: Counter) -> dict[str, Any]:
    mapped = {code: n for code, n in counts.items() if code in TAXONOMY_MAP}
    unmapped = {code: n for code, n in counts.items() if code not in TAXONOMY_MAP}
    not_in_map = {code: n for code, n in unmapped.items() if code not in UNMAPPED}
    return {
        "codes": len(counts),
        "rows": sum(counts.values()),
        "mapped_codes": len(mapped),
        "mapped_rows": sum(mapped.values()),
        "unmapped_codes": len(unmapped),
        "unmapped_rows": sum(unmapped.values()),
        "not_in_map_codes": len(not_in_map),
        "not_in_map_rows": sum(not_in_map.values()),
    }


def _pct(part: int, whole: int) -> str:
    return f"{100 * part / whole:.1f}%" if whole else "0.0%"


def summary_lines(summary: dict[str, Any], exceptions: Sequence[dict[str, Any]], top: int = 15) -> list[str]:
    total = summary["rows"]
    lines = [
        f"Input: {summary['codes']} distinct codes on {total} rows.",
        f"Mapped: {summary['mapped_codes']} codes, {summary['mapped_rows']} rows "
        f"({_pct(summary['mapped_rows'], total)}).",
        f"Unmapped: {summary['unmapped_codes']} codes, {summary['unmapped_rows']} rows "
        f"({_pct(summary['unmapped_rows'], total)}); {summary['not_in_map_codes']} of them "
        f"({summary['not_in_map_rows']} rows) in neither dict ({NOT_IN_MAP}).",
        "Top unmapped by rows:",
    ]
    lines += [f"  {row['code']}  {row['rows']:>6}  {row['reason']}" for row in exceptions[:top]]
    return lines


def report_lines(
    source: str,
    summary: dict[str, Any],
    exceptions: Sequence[dict[str, Any]],
    counts: Counter,
    groups: Sequence[dict[str, Any]],
    db: dict[str, Any],
    dry_run: bool,
    not_doctors: Counter | None = None,
) -> list[str]:
    wanted = desired_codes([group["slug"] for group in groups])
    outcome = "dry run (rolled back)" if dry_run else "committed"
    lines = [
        "## Taxonomy mapping (map_taxonomy.py)",
        "",
        f"Last run: {utc_now():%Y-%m-%d %H:%M} UTC, {outcome}. Input: {source}.",
        "",
    ]
    if not_doctors is not None:
        excluded = sum(not_doctors.values())
        lines += [
            "From the shortlist: the rows the registry import takes, i.e. the doctor sheets ("
            + ", ".join(f'"{sheet}"' for sheet in SHORTLIST_SHEETS)
            + f"), less {excluded} row(s) whose Profession and taxonomy code are not a doctor's"
            + (" (" + "; ".join(f"{p} {c}: {n}" for (p, c), n in sorted(not_doctors.items())) + ")" if excluded else "")
            + ". The other professions' codes keep their reasons in `scripts/taxonomy_map.py` UNMAPPED.",
            "",
        ]
    lines += markdown_table(
        ("", "Codes", "Rows", "Share of rows"),
        [
            ("Input", summary["codes"], summary["rows"], "100%"),
            (
                "Mapped to a group",
                summary["mapped_codes"],
                summary["mapped_rows"],
                _pct(summary["mapped_rows"], summary["rows"]),
            ),
            (
                "Unmapped (in output/taxonomy_exceptions.csv)",
                summary["unmapped_codes"],
                summary["unmapped_rows"],
                _pct(summary["unmapped_rows"], summary["rows"]),
            ),
            (
                f"of which {NOT_IN_MAP}",
                summary["not_in_map_codes"],
                summary["not_in_map_rows"],
                _pct(summary["not_in_map_rows"], summary["rows"]),
            ),
        ],
    )
    changed = db["changed"]
    lines += [
        "",
        f"Database: {len(changed)} group(s) changed; {db['groups_with_codes']} of the 104 groups hold codes; "
        f"{db['codes_written']} codes in place.",
    ]
    if changed:
        lines.append(
            "Changed: " + "; ".join(f"{c['slug']} (+{len(c['added'])} -{len(c['removed'])})" for c in changed) + "."
        )
    if db["missing_groups"]:
        missing = ", ".join(db["missing_groups"])
        lines.append(f"Groups not in the database (run seed_specialties.py), codes not written: {missing}.")
    lines += ["", "### Top unmapped codes by rows", ""]
    if not exceptions:
        lines.append("None: every taxonomy code of the input maps to a group (output/taxonomy_exceptions.csv holds the "
                     "header only).")
    else:
        lines += markdown_table(
            ("Code", "Rows", "Profession", "NUCC classification", "Reason"),
            [
                (
                    row["code"],
                    row["rows"],
                    row["profession"],
                    " / ".join(part for part in (row["classification"], row["specialization"]) if part),
                    row["reason"],
                )
                for row in exceptions[:20]
            ],
        )
    lines += [
        "",
        "### Codes and rows per group",
        "",
        "Codes are the group's codes in `scripts/taxonomy_map.py`; rows are registry rows in the input whose code maps to "
        "the group (from the shortlist: the doctors the import takes). A group with no codes gets none.",
        "",
    ]
    lines += markdown_table(
        ("#", "Group", "Codes", "Rows"),
        [
            (
                group["position"],
                f"`{group['slug']}`",
                len(wanted[group["slug"]]),
                sum(counts.get(code, 0) for code in wanted[group["slug"]]),
            )
            for group in groups
        ],
    )
    return lines


# ---------------------------------------------------------------- main


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument(
        "--shortlist", type=pathlib.Path,
        help="the NPPES shortlist workbook (.xlsx); its three doctor sheets are read, as the registry import reads them",
    )
    source.add_argument("--codes-csv", type=pathlib.Path, help="a CSV of codes with an optional count column")
    parser.add_argument(
        "--nucc", type=pathlib.Path, help="the NUCC taxonomy CSV, for the exceptions file's descriptions"
    )
    add_database_argument(parser)
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="write the exceptions file, compute the database changes, roll back; print the report section instead of "
        "writing it",
    )
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    groups = specialty_rows()
    slugs = [group["slug"] for group in groups]
    check_mapping(slugs)

    path = args.shortlist or args.codes_csv
    if not path.is_file():
        print(f"No such file: {path}", file=sys.stderr)
        return 2
    try:
        not_doctors: Counter | None = None
        if args.shortlist:
            counts, professions, not_doctors = read_shortlist(path)
            descriptions: dict[str, dict[str, str]] = {}
        else:
            counts, professions, descriptions = read_codes_csv(path)
        if args.nucc:
            if not args.nucc.is_file():
                print(f"No such file: {args.nucc}", file=sys.stderr)
                return 2
            descriptions = {**descriptions, **read_nucc(args.nucc)}
    except InputError as error:
        print(f"Cannot read the input: {error}", file=sys.stderr)
        return 2

    summary = summarise(counts)
    exceptions = build_exceptions(counts, professions, descriptions)
    exceptions_path = write_exceptions(exceptions)

    conn = connect(args.database_url)
    try:
        db = apply_to_database(conn, slugs)
        if args.dry_run:
            conn.rollback()
        else:
            conn.commit()
    except RuntimeError as error:
        conn.rollback()
        print(f"Stopped: {error}", file=sys.stderr)
        return 1
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    label = f"{'shortlist' if args.shortlist else 'codes CSV'} `{path.name}`" + (
        f", NUCC `{args.nucc.name}`" if args.nucc else ""
    )
    lines = report_lines(label, summary, exceptions, counts, groups, db, args.dry_run, not_doctors)
    report = None if args.dry_run else write_report_section("map_taxonomy", lines)

    if not_doctors is not None:
        print(f"Doctor sheets: {summary['rows']} rows the import takes; {sum(not_doctors.values())} not a doctor's"
              " (not counted).")
    for line in summary_lines(summary, exceptions):
        print(line)
    verb = "would change" if args.dry_run else "changed"
    print(
        f"Database: {len(db['changed'])} group(s) {verb}; "
        f"{db['codes_written']} codes across {db['groups_with_codes']} groups."
    )
    if db["missing_groups"]:
        print(
            f"Not in the database (run seed_specialties.py first): {len(db['missing_groups'])} groups, "
            f"{len(db['skipped_codes'])} codes skipped."
        )
    print(f"Exceptions: {exceptions_path.relative_to(exceptions_path.parents[1])} ({len(exceptions)} codes)")
    if report is None:
        print("\n".join(["", *lines, ""]))
        target = seed_specialties.REPORT_PATH
        print(f"Dry run: rolled back; {target.relative_to(target.parents[1]).as_posix()} not written.")
    else:
        print(f"Report: {report.relative_to(report.parents[1])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
